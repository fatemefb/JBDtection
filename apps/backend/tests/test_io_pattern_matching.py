"""Strict learned IO List matching and export regression coverage."""
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock

import fitz
import pandas as pd
from openpyxl import load_workbook

from jb_detection.annotator import PDFAnnotator, CATEGORY_COLORS_RGB
from jb_detection.config import Config
from jb_detection.excel_exporter import ExcelExporter
from jb_detection.models import OcrDetection
from jb_detection.pattern_matcher import PatternMatcher
from jb_detection.tag_matcher import TagMatcher
from jb_detection.unified_pdf_processor import UnifiedPdfProcessor


def det(text, y):
    return OcrDetection(text, 0.99, [], (20, y, 120, 20))


class IoPatternMatchingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.config = Config(pdf_dpi=72)
        self.io_path = self.root / 'io.xlsx'
        pd.DataFrame({'Tag No': ['TE-5223', '11-FV-301', '21HS-001'], 'Description': ['sensor', 'valve', 'switch']}).to_excel(self.io_path, index=False)
        self.matcher = TagMatcher(config=self.config)
        self.matcher.build_from_excel(str(self.io_path))

    def test_learns_numeric_and_area_families_but_rejects_unknown_prefixes(self):
        for tag in ['TE-5224', 'TE-522', 'TE-123456A', '12-FV-302', '123-FV-98765B', '22HS-002']:
            self.assertTrue(self.matcher.matches_io_pattern(tag), tag)
        for tag in ['PT-5223', '11-XV-301', '21ZZ-001', 'text TE-5223 extra', 'TE-5223ABC']:
            self.assertFalse(self.matcher.matches_io_pattern(tag), tag)
            self.assertEqual(self.matcher.match_tag(tag), ('unmatched', 0.0, ''))

    def test_vectors_only_rank_tags_in_the_learned_family(self):
        kind, score, closest = self.matcher.match_tag('TE-5224')
        self.assertEqual(kind, 'similar')
        self.assertGreater(score, 0.85)
        self.assertLess(score, 1)
        self.assertEqual(closest, 'TE-5223')
        self.assertEqual(self.matcher.match_tag('TE-5223'), ('exact', 1.0, 'TE-5223'))

    def test_rebuilding_clears_patterns_and_vectors_from_previous_list(self):
        pd.DataFrame({'Tag No': ['PT-1000']}).to_excel(self.io_path, index=False)
        self.matcher.build_from_excel(str(self.io_path))
        self.assertFalse(self.matcher.matches_io_pattern('TE-5224'))
        self.assertEqual(set(self.matcher.tag_vectors), {'PT-1000'})

    def test_pattern_gate_controls_classification_and_nonstandard_tags(self):
        pattern = PatternMatcher()
        pattern.io_tag_matcher = self.matcher
        result = pattern.match([det('PT-5223', 10), det('12-FV-302', 30), det('22HS-002', 50), det('TE-5224ABC', 70)])
        self.assertEqual(result.tags, {'12-FV-302', '22HS-002'})
        self.assertEqual(result.all_ocr_tags, result.tags)
        self.assertEqual(set(result.tag_to_number), result.tags)

    def test_empty_io_list_fails_closed(self):
        pattern = PatternMatcher()
        pattern.io_tag_matcher = TagMatcher()
        self.assertEqual(pattern.match([det('TE-5223', 10)]).tags, set())

    def results(self):
        proc = UnifiedPdfProcessor(config=self.config, tag_matcher=self.matcher)
        return proc._process_detections([
            det('JB-101', 10), det('TE-5223', 40), det('TE-5224', 80),
            det('TE-9999', 120), det('PT-5223', 160),
        ], 1)

    def test_pdf_exact_similar_and_absent_tags_have_distinct_colors(self):
        result = self.results()
        path = self.root / 'blank.pdf'
        with fitz.open() as doc:
            doc.new_page()
            doc.save(path)
        PDFAnnotator(config=self.config).annotate_pdf(str(path), {1: result}, str(self.root / 'annotated.pdf'), tag_to_number=result.tag_to_number)
        with fitz.open(self.root / 'annotated.pdf') as doc:
            colors = {tuple(round(value, 2) for value in drawing['color']) for drawing in doc[0].get_drawings()}
            for category in ['tag', 'mc', 'unknown']:
                self.assertIn(tuple(round(value, 2) for value in CATEGORY_COLORS_RGB[category]), colors)
            self.assertIn('[S:', doc[0].get_text())
            self.assertNotIn('PT-5223', doc[0].get_text())

    def test_excel_preserves_absent_tag_wiring_without_replacing_reference(self):
        result = self.results()
        exporter = ExcelExporter(config=self.config)
        pattern = PatternMatcher()
        pattern.io_tag_matcher = self.matcher
        exporter.set_pattern_matcher(pattern)
        intermediate = self.root / 'intermediate.xlsx'
        exporter.create_intermediate_excel({'drawing.pdf': {1: result}}, str(intermediate))
        output = self.root / 'final.xlsx'
        frame, missing_io, missing_pdf = exporter.create_final_excel(str(intermediate), str(self.io_path), str(output), result.all_ocr_tags)
        indexed = frame.set_index('Tag No')
        self.assertEqual(set(indexed.index), {'TE-5223', '11-FV-301', '21HS-001', 'TE-5224', 'TE-9999'})
        self.assertEqual(indexed.loc['TE-5223', 'Match_Type'], 'exact')
        candidate = indexed.loc['TE-5224']
        self.assertEqual(candidate['Match_Type'], 'similar')
        self.assertEqual(candidate['Closest_IO_Tag'], 'TE-5223')
        self.assertEqual(candidate['JB'], 'JB-101')
        self.assertEqual(candidate['Tag_Number'], 2)
        self.assertTrue(candidate['Terminal_First_Number'])
        self.assertTrue(candidate['Wire_Code_1'])
        self.assertTrue(pd.isna(candidate['Description']))
        self.assertEqual(set(missing_pdf), {'TE-5224', 'TE-9999'})
        self.assertIn('11-FV-301', missing_io)
        self.assertEqual(len(exporter._last_pattern_candidates), 2)
        detail = next(item for item in exporter._last_pattern_candidates if item['ocr_text'] == 'TE-5224')
        self.assertEqual(detail['tag_number'], 2)
        self.assertEqual(detail['closest_io_tag'], 'TE-5223')
        self.assertTrue(detail['terminal_first_number'])
        book = load_workbook(output)
        rows = list(book.active.values)
        kind_index = rows[0].index('Match_Type')
        for index, row in enumerate(rows[1:], start=2):
            if row[kind_index] == 'similar':
                self.assertEqual(book.active.cell(index, 1).fill.fgColor.rgb, '00FFF2CC')
        book.close()

    def test_full_compat_entrypoint_exposes_candidates_and_exports_original_tag(self):
        from jb_detection.compat import TagJBExtractor
        pdf = self.root / 'drawing.pdf'
        with fitz.open() as doc:
            page = doc.new_page()
            for tag, y in [('JB-101', 50), ('TE-5224', 100), ('PT-5223', 150), ('11-FV-301', 200)]:
                page.insert_text((50, y), tag, fontsize=14)
            doc.save(pdf)
        extractor = TagJBExtractor()
        extractor._processor._detector = MagicMock()
        output = self.root / 'final.xlsx'
        missing_io, missing_pdf = extractor.run_with_annotated_pdf(
            [str(pdf)], str(self.io_path), str(output), str(self.root / 'annotated'), create_zip=False,
        )
        extractor._processor._detector.detect.assert_not_called()
        self.assertEqual(extractor.latest_pattern_unmatched_candidates, ['TE-5224'])
        detail = extractor.latest_pattern_unmatched_details[0]
        self.assertEqual(detail['closest_io_tag'], 'TE-5223')
        self.assertEqual(detail['tag_number'], 1)
        self.assertEqual(detail['jb'], 'JB-101')
        self.assertTrue(detail['terminal_first_number'])
        self.assertIn('TE-5223', missing_io)
        self.assertIn('TE-5224', missing_pdf)
        self.assertNotIn('PT-5223', missing_pdf)
        self.assertIn('TE-5224', pd.read_excel(output)['Tag No'].tolist())

    def test_configured_similarity_threshold_and_batch_api(self):
        matcher = TagMatcher(config=Config(match_similar_threshold=0.99))
        matcher.build_from_excel(str(self.io_path))
        matches = matcher.match_many(['TE-5223', 'TE-5224', 'PT-5223'])
        self.assertEqual(matches['TE-5223'][0], 'exact')
        self.assertEqual(matches['TE-5224'][0], 'unmatched_candidate')
        self.assertEqual(matches['TE-5224'][2], 'TE-5223')
        self.assertEqual(matches['PT-5223'], ('unmatched', 0.0, ''))

    def test_absent_candidates_survive_without_a_similarity_match(self):
        # An intentionally unreachable threshold proves discovery does not
        # depend on fuzzy acceptance, even for new serial widths/suffixes.
        self.matcher.similarity_threshold = 1.0
        proc = UnifiedPdfProcessor(config=self.config, tag_matcher=self.matcher)
        result = proc._process_detections([
            det('TE-987654A', 10), det('123-FV-98765B', 40),
            det('TE-777777 TE-888888A', 70), det('ZZ-987654A', 100),
        ], 1)
        expected = {'TE-987654A', '123-FV-98765B', 'TE-777777', 'TE-888888A'}
        self.assertEqual(result.tags, expected)
        self.assertEqual(result.all_ocr_tags, expected)
        self.assertEqual(set(result.tag_to_number), expected)
        for tag in expected:
            self.assertEqual(result.tag_match_info[tag].match_type, 'unmatched_candidate')
            self.assertIn('absent', result.tag_match_info[tag].reason)
        self.assertEqual(result.tag_match_info['TE-987654A'].bbox, (20, 10, 120, 20))

    def test_legacy_terminal_letter_variants_remain_candidates(self):
        matcher = TagMatcher()
        matcher.add_reference_tag('LUSY-2474A')
        matcher.add_reference_tag('USY-2482A')
        candidates = matcher.extract_candidates('LUSY-99999B USY-1 USY-876543C XYZ-2474A')
        self.assertEqual(candidates, ['LUSY-99999B', 'USY-1', 'USY-876543C'])
        for tag in candidates:
            self.assertEqual(matcher.match_tag(tag)[0], 'unmatched_candidate')
        self.assertEqual(matcher.extract_candidates('X.LUSY-99999B LUSY-99999B.1'), [])

    def test_low_similarity_candidate_keeps_exact_text_in_excel_and_ui_metadata(self):
        proc = UnifiedPdfProcessor(config=self.config, tag_matcher=self.matcher)
        result = proc._process_detections([det('JB-101', 10), det('TE-987654A', 40)], 1)
        exporter = ExcelExporter(config=self.config)
        exporter.set_pattern_matcher(proc.pattern_matcher)
        intermediate = self.root / 'candidate-intermediate.xlsx'
        exporter.create_intermediate_excel({'candidate.pdf': {1: result}}, str(intermediate))
        final = self.root / 'candidate-final.xlsx'
        frame, missing_io, missing_pdf = exporter.create_final_excel(str(intermediate), str(self.io_path), str(final), result.all_ocr_tags)
        candidate = frame[frame['Tag No'] == 'TE-987654A'].iloc[0]
        self.assertEqual(candidate['Match_Type'], 'unmatched_candidate')
        self.assertEqual(candidate['Tag_Number'], 1)
        self.assertEqual(candidate['JB'], 'JB-101')
        self.assertIn('TE-987654A', missing_pdf)
        self.assertIn('TE-5223', missing_io)
        detail = exporter._last_pattern_candidates[0]
        self.assertEqual(detail['ocr_text'], 'TE-987654A')
        self.assertEqual(detail['match_type'], 'unmatched_candidate')
        self.assertTrue(detail['terminal_first_number'])
        self.assertTrue(detail['wire_code_1'])
        book = load_workbook(final)
        tag_index = list(next(book.active.values)).index('Tag No')
        row_index = next(index for index, row in enumerate(list(book.active.values)[1:], start=2) if row[tag_index] == 'TE-987654A')
        self.assertEqual(book.active.cell(row_index, 1).fill.fgColor.rgb, '00FCE4D6')
        book.close()

    def test_native_span_candidates_have_individual_word_boxes(self):
        pdf = self.root / 'one-span.pdf'
        with fitz.open() as doc:
            page = doc.new_page()
            page.insert_text((30, 60), 'JB-101 TE-987654A TE-888888A PT-987654A', fontsize=12)
            doc.save(pdf)
        detector = MagicMock()
        proc = UnifiedPdfProcessor(config=self.config, tag_matcher=self.matcher, detector=detector)
        result = proc.process_pdf(str(pdf))[1]
        detector.detect.assert_not_called()
        self.assertEqual(result.tags, {'TE-987654A', 'TE-888888A'})
        self.assertEqual(result.jb_identifiers, {'JB-101'})
        self.assertEqual(proc.pages_digital, 1)
        boxes = [result.tag_match_info[tag].bbox for tag in sorted(result.tags)]
        self.assertNotEqual(boxes[0], boxes[1])
        self.assertTrue(boxes[0][0] + boxes[0][2] < boxes[1][0] or boxes[1][0] + boxes[1][2] < boxes[0][0])
        for tag in result.tags:
            self.assertEqual(result.tag_match_info[tag].match_type, 'unmatched_candidate')


if __name__ == '__main__':
    unittest.main()
